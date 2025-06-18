import os
import json
import exifread
import PyPDF2
import docx
import mutagen
from PIL import Image
from hachoir.parser import createParser
from hachoir.metadata import extractMetadata
import google.generativeai as genai
from dotenv import load_dotenv

# Load environment variables
load_dotenv()

# Configure the Gemini API with the API key
GEMINI_API_KEY = os.getenv('GEMINI_API_KEY')
if GEMINI_API_KEY:
    genai.configure(api_key=GEMINI_API_KEY)
else:
    print("GEMINI_API_KEY not found in environment variables. Please set it in .env file.")
    exit(1)

def extract_image_metadata(file_path):
    """Extract metadata from image files"""
    metadata = {}

    # Extract EXIF data using exifread
    with open(file_path, 'rb') as f:
        tags = exifread.process_file(f)
        for tag, value in tags.items():
            metadata[tag] = str(value)

    # Extract basic image info using PIL
    try:
        with Image.open(file_path) as img:
            metadata['format'] = img.format
            metadata['mode'] = img.mode
            metadata['size'] = str(img.size)
            if hasattr(img, 'info'):
                for key, value in img.info.items():
                    if isinstance(value, (str, int, float, bool)):
                        metadata[f'info_{key}'] = value
    except Exception as e:
        metadata['pil_error'] = str(e)

    return metadata

def extract_pdf_metadata(file_path):
    """Extract metadata from PDF files"""
    metadata = {}

    try:
        with open(file_path, 'rb') as f:
            pdf = PyPDF2.PdfReader(f)
            if pdf.metadata:
                for key, value in pdf.metadata.items():
                    metadata[key] = str(value)
            metadata['pages'] = len(pdf.pages)
    except Exception as e:
        metadata['error'] = str(e)

    return metadata

def extract_docx_metadata(file_path):
    """Extract metadata from DOCX files"""
    metadata = {}

    try:
        doc = docx.Document(file_path)
        core_properties = doc.core_properties
        metadata['author'] = core_properties.author
        metadata['created'] = str(core_properties.created)
        metadata['modified'] = str(core_properties.modified)
        metadata['title'] = core_properties.title
        metadata['subject'] = core_properties.subject
        metadata['keywords'] = core_properties.keywords
        metadata['comments'] = core_properties.comments
        metadata['category'] = core_properties.category
        metadata['content_status'] = core_properties.content_status
        metadata['paragraph_count'] = len(doc.paragraphs)
    except Exception as e:
        metadata['error'] = str(e)

    return metadata

def extract_audio_metadata(file_path):
    """Extract metadata from audio files"""
    metadata = {}

    try:
        audio = mutagen.File(file_path)
        if audio:
            for key, value in audio.items():
                metadata[key] = str(value)
            if hasattr(audio, 'info'):
                metadata['length'] = audio.info.length
                metadata['bitrate'] = audio.info.bitrate
    except Exception as e:
        metadata['error'] = str(e)

    return metadata

def extract_hachoir_metadata(file_path):
    """Extract metadata using hachoir (works for many file types)"""
    metadata = {}

    try:
        parser = createParser(file_path)
        if parser:
            metadata_obj = extractMetadata(parser)
            if metadata_obj:
                for line in metadata_obj.exportPlaintext():
                    if ': ' in line:
                        key, value = line.split(': ', 1)
                        metadata[key.strip()] = value.strip()
    except Exception as e:
        metadata['hachoir_error'] = str(e)

    return metadata

def analyze_metadata_with_gemini(metadata_dict, file_info):
    """Analyze metadata using Google's Gemini API"""
    try:
        # Format the metadata for better readability
        formatted_metadata = json.dumps(metadata_dict, indent=2)

        # Create a prompt for Gemini
        prompt = f"""
        I have extracted metadata from a file with the following information:

        File name: {file_info.get('filename', 'Unknown')}
        File type: {file_info.get('file_type', 'Unknown')}

        Here is the extracted metadata:

        {formatted_metadata}

        Please analyze this metadata and provide the following:

        1. Anomalies: Does this metadata contain any anomalies or signs of tampering? If so, what are they?
        2. Privacy Concerns: Does this metadata contain any privacy-sensitive information that the user should be aware of?
        3. Summary: Provide a concise summary of the most important metadata information.

        Format your response as JSON. Use the following structure, replacing the placeholders with your analysis:
        {{"anomalies":{{"detected":false,"details":"No anomalies detected"}},"privacy_concerns":{{"detected":false,"details":"No privacy concerns detected"}},"summary":"This is where you put your summary"}}
        """

        # Generate content using Gemini
        model = genai.GenerativeModel('gemini-1.5-pro')
        response = model.generate_content(prompt)

        # Parse the response
        try:
            # Try to extract JSON from the response
            response_text = response.text
            # Find JSON content (it might be wrapped in markdown code blocks)
            if "```json" in response_text:
                json_content = response_text.split("```json")[1].split("```")[0].strip()
            elif "```" in response_text:
                json_content = response_text.split("```")[1].split("```")[0].strip()
            else:
                json_content = response_text

            analysis_result = json.loads(json_content)
            return analysis_result

        except json.JSONDecodeError:
            # If JSON parsing fails, return the raw text
            return {
                "error": "Failed to parse JSON response",
                "raw_response": response.text
            }

    except Exception as e:
        return {
            "error": str(e)
        }

def test_metadata_extraction(file_path):
    """Test metadata extraction and analysis on a file"""
    if not os.path.exists(file_path):
        print(f"File not found: {file_path}")
        return

    print(f"Extracting metadata from: {file_path}")

    # Get file info
    file_info = {
        'filename': os.path.basename(file_path),
        'file_type': os.path.splitext(file_path)[1].lower(),
        'file_size': os.path.getsize(file_path)
    }

    # Extract metadata based on file type
    metadata = {}
    file_type = file_info['file_type']

    if file_type in ['.jpg', '.jpeg', '.png', '.gif', '.bmp', '.tiff']:
        metadata = extract_image_metadata(file_path)
    elif file_type == '.pdf':
        metadata = extract_pdf_metadata(file_path)
    elif file_type == '.docx':
        metadata = extract_docx_metadata(file_path)
    elif file_type in ['.mp3', '.wav', '.flac', '.ogg', '.m4a']:
        metadata = extract_audio_metadata(file_path)

    # Always try hachoir as a fallback
    hachoir_metadata = extract_hachoir_metadata(file_path)
    metadata.update(hachoir_metadata)

    # Print extracted metadata
    print("\nExtracted Metadata:")
    print(json.dumps(metadata, indent=2))

    # Analyze metadata with Gemini
    print("\nAnalyzing metadata with Gemini API...")
    analysis = analyze_metadata_with_gemini(metadata, file_info)

    # Print analysis results
    print("\nAnalysis Results:")
    print(json.dumps(analysis, indent=2))

if __name__ == "__main__":
    # Test with a sample file - replace with your own file path
    import sys
    if len(sys.argv) > 1:
        test_file = sys.argv[1]
    else:
        print("Please provide a file path as an argument")
        print("Example: python test_metadata.py path/to/your/file.jpg")
        sys.exit(1)

    test_metadata_extraction(test_file)
